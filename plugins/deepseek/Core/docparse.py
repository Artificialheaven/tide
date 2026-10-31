"""群友发来的附件解析为可送模文本（纯文本 / PDF / Word / Excel）。

- 纯文本类直接解码（utf-8-sig → utf-8 → gbk 兜底，Windows 常见 GBK txt）
- PDF 走 pypdf、docx 走 python-docx、xlsx 走 openpyxl、老 xls 走 xlrd：
  全部延迟导入，缺依赖时返回一句说明而不是报错
- .doc（老二进制 Word）不支持，提示另存 docx
解析在 to_thread 里跑（大文件 CPU/IO 都不阻塞事件循环）。
"""
from __future__ import annotations

import io

_TEXT_EXTS = {
    'txt', 'md', 'markdown', 'log', 'csv', 'tsv', 'json', 'xml', 'yaml', 'yml',
    'ini', 'conf', 'py', 'js', 'ts', 'java', 'c', 'h', 'cpp', 'cc', 'hpp',
    'cs', 'go', 'rs', 'rb', 'php', 'sh', 'bat', 'ps1', 'sql', 'html', 'htm',
    'css', 'vue', 'kt', 'swift', 'lua', 'toml', 'properties', 'gitignore',
}


def _decode(raw: bytes) -> str:
    for enc in ('utf-8-sig', 'utf-8', 'gbk'):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode('utf-8', errors='replace')


def _parse_pdf(raw: bytes) -> tuple[str | None, str]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return None, 'PDF 解析依赖 pypdf（未安装），无法读取内容'
    try:
        reader = PdfReader(io.BytesIO(raw))
        parts = []
        for page in reader.pages[:50]:
            parts.append(page.extract_text() or '')
        text = '\n'.join(parts).strip()
        note = '' if text else 'PDF 中没有提取到文本（可能是扫描件图片）'
        return text, note
    except Exception as e:
        return None, f'PDF 解析失败: {e}'


def _parse_docx(raw: bytes) -> tuple[str | None, str]:
    try:
        import docx
    except ImportError:
        return None, 'Word 解析依赖 python-docx（未安装），无法读取内容'
    try:
        doc = docx.Document(io.BytesIO(raw))
        parts = [p.text for p in doc.paragraphs if p.text]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells]
                if any(cells):
                    parts.append(' | '.join(cells))
        text = '\n'.join(parts).strip()
        note = '' if text else 'Word 文档里没有提取到文本'
        return text, note
    except Exception as e:
        return None, f'Word 解析失败: {e}'


def _render_sheets(sheets: list) -> str:
    """sheets: [(表名, [[行单元格...], ...]), ...] → 管道分隔文本。"""
    parts = []
    for name, rows in sheets:
        lines = []
        for row in rows:
            cells = ['' if v is None else str(v).strip() for v in row]
            if any(cells):
                # 去掉行尾空单元格，减少噪声
                while cells and not cells[-1]:
                    cells.pop()
                lines.append(' | '.join(cells))
        if lines:
            parts.append(f'【工作表 {name}】\n' + '\n'.join(lines))
    return '\n'.join(parts).strip()


def _parse_xlsx(raw: bytes) -> tuple[str | None, str]:
    try:
        import openpyxl
    except ImportError:
        return None, 'Excel 解析依赖 openpyxl（未安装），无法读取内容'
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True,
                                    data_only=True)
        sheets = []
        for ws in wb.worksheets[:10]:
            rows = []
            for row in ws.iter_rows(values_only=True):
                rows.append(row)
                if len(rows) >= 2000:
                    break
            sheets.append((ws.title, rows))
        text = _render_sheets(sheets)
        note = '' if text else 'Excel 工作簿里没有提取到文本'
        return text, note
    except Exception as e:
        return None, f'Excel 解析失败: {e}'


def _parse_xls(raw: bytes) -> tuple[str | None, str]:
    try:
        import xlrd
    except ImportError:
        return None, '老版 .xls 解析依赖 xlrd（未安装），请另存为 .xlsx 后重发'
    try:
        wb = xlrd.open_workbook(file_contents=raw)
        sheets = []
        for sh in wb.sheets()[:10]:
            rows = [[sh.cell_value(r, c) for c in range(sh.ncols)]
                    for r in range(min(sh.nrows, 2000))]
            sheets.append((sh.name, rows))
        text = _render_sheets(sheets)
        note = '' if text else 'Excel 工作簿里没有提取到文本'
        return text, note
    except Exception as e:
        return None, f'Excel 解析失败: {e}'


def extract_text(filename: str, raw: bytes, max_chars: int = 4000) -> str:
    """把附件解析成给模型看的文本块；不支持/失败时返回一句方括号说明。"""
    name = (filename or '附件')
    ext = name.rsplit('.', 1)[-1].lower() if '.' in name else ''
    size = len(raw)
    head = f'[附件 {name}（{size} 字节）'
    if ext in _TEXT_EXTS:
        body, note = _decode(raw), ''
    elif ext == 'pdf':
        body, note = _parse_pdf(raw)
    elif ext == 'docx':
        body, note = _parse_docx(raw)
    elif ext in ('xlsx', 'xlsm'):
        body, note = _parse_xlsx(raw)
    elif ext == 'xls':
        body, note = _parse_xls(raw)
    elif ext == 'doc':
        body, note = None, '老版 .doc 不支持，请另存为 .docx 后重发'
    else:
        return f'{head}：未支持的文件类型 .{ext or "?"}]'
    if body is None:
        return f'{head}：{note}]'
    body = body.strip()
    if not body:
        return f'{head}：{note or "空文件"}]'
    if len(body) > max_chars:
        body = body[:max_chars] + f'\n…（已截断，全文 {len(body)} 字符）'
    return f'{head} 的文本内容]\n{body}'
